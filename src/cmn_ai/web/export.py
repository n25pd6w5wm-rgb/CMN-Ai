"""Render an assistant answer (Markdown) into polished documents (PDF/DOCX/PPTX).

The parser understands the Markdown the models emit — headings, ordered and
unordered (nestable) lists, fenced code, pipe tables, blockquotes, horizontal
rules, and inline **bold** / *italic* / ``code`` / [links] — and each renderer
turns it into a properly typeset document.

Several visual **themes** are available so generated documents don't all look the
same: ``report`` (formal blue, the default), ``modern`` (teal, with a cover page),
``elegant`` (restrained terracotta, cover page) and ``deck`` (indigo, the default
for presentations). The model can request one via the ``cmn:file`` fence; otherwise
a sensible per-format default is used.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

_FONT_DIR = Path(__file__).resolve().parent / "fonts"

RGB = tuple[int, int, int]


# ---------- themes (documents are always on white, independent of the app theme) ----------


@dataclass(frozen=True)
class Theme:
    """A visual style shared across the PDF, DOCX and PPTX renderers."""

    name: str
    accent: RGB  # headings, bars, rules, links
    accent_dark: RGB  # cover band / darker accent for contrast
    head_bg: RGB  # table-header + soft fills
    ink: RGB  # body text
    muted: RGB  # captions, footers, quotes
    code_bg: RGB  # code block background
    rule: RGB  # horizontal rules / hairlines
    cover: bool = False  # render a dedicated cover page (PDF) / title slide accent (PPTX)
    title_size: float = 24.0  # cover / document title size
    heading_accent: bool = True  # colour H1/H2 in the accent (False = restrained, ink headings)


THEMES: dict[str, Theme] = {
    "report": Theme(
        name="report",
        accent=(43, 127, 196),
        accent_dark=(28, 86, 140),
        head_bg=(237, 242, 248),
        ink=(28, 28, 30),
        muted=(110, 108, 104),
        code_bg=(244, 245, 247),
        rule=(222, 222, 220),
        cover=False,
        title_size=24.0,
        heading_accent=True,
    ),
    "modern": Theme(
        name="modern",
        accent=(13, 148, 136),
        accent_dark=(11, 94, 90),
        head_bg=(224, 242, 239),
        ink=(23, 37, 42),
        muted=(90, 110, 110),
        code_bg=(240, 247, 246),
        rule=(210, 226, 223),
        cover=True,
        title_size=30.0,
        heading_accent=True,
    ),
    "elegant": Theme(
        name="elegant",
        accent=(161, 66, 52),
        accent_dark=(112, 43, 34),
        head_bg=(245, 236, 232),
        ink=(38, 34, 32),
        muted=(120, 105, 98),
        code_bg=(247, 243, 240),
        rule=(226, 216, 210),
        cover=True,
        title_size=30.0,
        heading_accent=False,
    ),
    "deck": Theme(
        name="deck",
        accent=(79, 70, 229),
        accent_dark=(55, 48, 163),
        head_bg=(232, 230, 252),
        ink=(24, 24, 37),
        muted=(100, 100, 120),
        code_bg=(243, 243, 250),
        rule=(220, 218, 240),
        cover=True,
        title_size=32.0,
        heading_accent=True,
    ),
    "minimal": Theme(
        name="minimal",
        accent=(26, 26, 26),
        accent_dark=(0, 0, 0),
        head_bg=(238, 238, 238),
        ink=(26, 26, 26),
        muted=(120, 120, 120),
        code_bg=(245, 245, 245),
        rule=(220, 220, 220),
        cover=False,
        title_size=26.0,
        heading_accent=False,
    ),
    "warm": Theme(
        name="warm",
        accent=(146, 100, 62),
        accent_dark=(92, 62, 38),
        head_bg=(244, 239, 230),
        ink=(28, 26, 24),
        muted=(125, 112, 100),
        code_bg=(247, 243, 237),
        rule=(228, 220, 210),
        cover=True,
        title_size=30.0,
        heading_accent=True,
    ),
}

# When the model doesn't name a theme, pick one that fits the format: documents get
# the formal report look; slide decks get the presentation-oriented deck theme.
_DEFAULT_THEME_FOR = {"pdf": "report", "docx": "report", "pptx": "deck", "xlsx": "report"}


def resolve_theme(name: str, fmt: str) -> Theme:
    """Resolve a (possibly empty or unknown) theme name for a format to a Theme."""
    if name and name in THEMES:
        return THEMES[name]
    return THEMES[_DEFAULT_THEME_FOR.get(fmt, "report")]


# ---------- inline parsing (bold / italic / code / links) ----------

_TOKEN = re.compile(
    r"(\*\*.+?\*\*|__.+?__|\*[^*]+?\*|_[^_]+?_|`[^`]+`|\[[^\]]+\]\([^)]+\))",
    re.DOTALL,
)


def _runs(text: str) -> list[tuple[str, frozenset[str]]]:
    """Split text into styled runs: (text, {"b"?, "i"?, "code"?})."""
    out: list[tuple[str, frozenset[str]]] = []
    pos = 0
    for m in _TOKEN.finditer(text):
        if m.start() > pos:
            out.append((text[pos : m.start()], frozenset()))
        tok = m.group(0)
        if tok.startswith(("**", "__")):
            out.append((tok[2:-2], frozenset({"b"})))
        elif tok.startswith("`"):
            out.append((tok[1:-1], frozenset({"code"})))
        elif tok.startswith("["):
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            label, url = (lm.group(1), lm.group(2)) if lm else (tok, "")
            out.append((f"{label} ({url})" if url else label, frozenset()))
        else:
            out.append((tok[1:-1], frozenset({"i"})))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], frozenset()))
    return out or [(text, frozenset())]


def _plain(text: str) -> str:
    return "".join(t for t, _ in _runs(text))


# ---------- block parsing ----------


@dataclass
class _Block:
    kind: str  # h1|h2|h3|h4 | para | bullet | number | code | table | quote | hr
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    level: int = 0
    marker: str = ""


_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBER = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_HEAD = re.compile(r"^(#{1,4})\s+(.*)$")


def _table_cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _is_sep_row(line: str) -> bool:
    body = line.strip().strip("|")
    return bool(body) and all(set(c.strip()) <= {"-", ":"} for c in body.split("|"))


def _parse(content: str) -> list[_Block]:
    blocks: list[_Block] = []
    lines = content.splitlines()
    i = 0
    para: list[str] = []

    def flush_para() -> None:
        if para:
            blocks.append(_Block("para", " ".join(para).strip()))
            para.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_para()
            i += 1
            continue

        # fenced code
        if stripped.startswith("```"):
            flush_para()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(_Block("code", "\n".join(code)))
            continue

        # pipe table
        if stripped.startswith("|") and "|" in stripped[1:]:
            flush_para()
            rows = [_table_cells(line)]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _is_sep_row(lines[i]):
                    rows.append(_table_cells(lines[i]))
                i += 1
            blocks.append(_Block("table", rows=rows))
            continue

        # horizontal rule
        if stripped in ("---", "***", "___"):
            flush_para()
            blocks.append(_Block("hr"))
            i += 1
            continue

        h = _HEAD.match(line)
        if h:
            flush_para()
            blocks.append(_Block(f"h{len(h.group(1))}", h.group(2).strip()))
            i += 1
            continue

        q = _QUOTE.match(line)
        if q:
            flush_para()
            quote = [q.group(1)]
            i += 1
            while i < len(lines) and (qm := _QUOTE.match(lines[i])):
                quote.append(qm.group(1))
                i += 1
            blocks.append(_Block("quote", " ".join(quote).strip()))
            continue

        n = _NUMBER.match(line)
        if n:
            flush_para()
            blocks.append(
                _Block("number", n.group(3).strip(), level=len(n.group(1)) // 2, marker=n.group(2))
            )
            i += 1
            continue

        b = _BULLET.match(line)
        if b:
            flush_para()
            blocks.append(_Block("bullet", b.group(2).strip(), level=len(b.group(1)) // 2))
            i += 1
            continue

        para.append(stripped)
        i += 1

    flush_para()
    return blocks


# ---------- PDF (fpdf2, vendored DejaVu, themed, page numbers) ----------


def _pdf_cover(pdf: Any, th: Theme, title: str) -> None:
    """A dedicated title page: accent band, big title, accent rule, byline."""
    pdf.add_page()
    pdf.set_fill_color(*th.accent)
    pdf.rect(0, 0, pdf.w, 60, style="F")
    pdf.set_xy(pdf.l_margin, 84)
    pdf.set_font("DejaVu", "B", th.title_size)
    pdf.set_text_color(*th.ink)
    pdf.multi_cell(0, th.title_size * 0.5, title, new_x="LMARGIN", new_y="NEXT")
    y = pdf.get_y() + 4
    pdf.set_draw_color(*th.accent)
    pdf.set_line_width(1.4)
    pdf.line(pdf.l_margin, y, pdf.l_margin + 55, y)
    pdf.set_xy(pdf.l_margin, y + 9)
    pdf.set_font("DejaVu", "", 12)
    pdf.set_text_color(*th.muted)
    pdf.cell(0, 8, "erstellt mit cmn·ai")


def to_pdf(content: str, *, theme: str = "") -> bytes:
    from fpdf import FPDF
    from fpdf.enums import TableBordersLayout
    from fpdf.fonts import FontFace

    th = resolve_theme(theme, "pdf")

    class _PDF(FPDF):
        muted: RGB = th.muted

        def footer(self) -> None:
            self.set_y(-14)
            self.set_font("DejaVu", "", 8)
            self.set_text_color(*self.muted)
            self.cell(0, 8, f"Seite {self.page_no()}", align="C")

    pdf = _PDF()
    for style, fname in (
        ("", "DejaVuSans.ttf"),
        ("B", "DejaVuSans-Bold.ttf"),
        ("I", "DejaVuSans-Oblique.ttf"),
        ("BI", "DejaVuSans-BoldOblique.ttf"),
    ):
        pdf.add_font("DejaVu", style, str(_FONT_DIR / fname))
    pdf.add_font("DejaVuMono", "", str(_FONT_DIR / "DejaVuSansMono.ttf"))
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(auto=True, margin=18)

    blocks = _parse(content)
    title_text = next((b.text for b in blocks if b.kind == "h1"), "")
    title_done = False
    if th.cover and title_text:
        _pdf_cover(pdf, th, _plain(title_text))
        title_done = True
    pdf.add_page()

    head_color = th.accent if th.heading_accent else th.ink

    def runs(text: str, size: float, base: str = "", color: RGB | None = None) -> None:
        color = th.ink if color is None else color
        lh = size * 0.52
        for txt, st in _runs(text):
            if "code" in st:
                pdf.set_font("DejaVuMono", "", size - 1)
                pdf.set_text_color(*th.accent)
            else:
                style = base + ("B" if "b" in st else "") + ("I" if "i" in st else "")
                pdf.set_font("DejaVu", "".join(sorted(style)), size)
                pdf.set_text_color(*color)
            pdf.write(lh, txt)
        pdf.ln(lh)

    for blk in blocks:
        if blk.kind == "h1" and not title_done:
            title_done = True
            pdf.set_font("DejaVu", "B", 24)
            pdf.set_text_color(*th.ink)
            pdf.multi_cell(0, 11, _plain(blk.text), new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(*th.accent)
            pdf.set_line_width(0.6)
            y = pdf.get_y() + 1
            pdf.line(pdf.l_margin, y, pdf.l_margin + 40, y)
            pdf.ln(6)
        elif blk.kind in ("h1", "h2", "h3", "h4"):
            size = {"h1": 17.0, "h2": 14.0, "h3": 12.0, "h4": 11.0}[blk.kind]
            pdf.ln(2)
            pdf.set_font("DejaVu", "B", size)
            pdf.set_text_color(*(head_color if blk.kind in ("h1", "h2") else th.ink))
            pdf.multi_cell(0, size * 0.55, _plain(blk.text), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
        elif blk.kind == "para":
            runs(blk.text, 11)
            pdf.ln(1.5)
        elif blk.kind in ("bullet", "number"):
            marker = f"{blk.marker}." if blk.kind == "number" else "•"
            indent = blk.level * 7
            left = pdf.l_margin + indent
            pdf.set_left_margin(left + 7)
            pdf.set_x(left)
            pdf.set_font("DejaVu", "", 11)
            pdf.set_text_color(*(th.accent if blk.kind == "number" else th.muted))
            pdf.write(11 * 0.52, f"{marker} ")
            pdf.set_x(left + 7)
            runs(blk.text, 11)
            pdf.set_left_margin(20)
        elif blk.kind == "quote":
            pdf.ln(1)
            top = pdf.get_y()
            pdf.set_left_margin(26)
            pdf.set_x(26)
            runs(blk.text, 11, base="I", color=th.muted)
            pdf.set_draw_color(*th.accent)
            pdf.set_line_width(1.2)
            pdf.line(21, top, 21, pdf.get_y() - 2)
            pdf.set_left_margin(20)
            pdf.ln(2)
        elif blk.kind == "code":
            pdf.ln(1)
            pdf.set_font("DejaVuMono", "", 9)
            pdf.set_text_color(*th.ink)
            pdf.set_fill_color(*th.code_bg)
            pdf.multi_cell(
                0, 5, blk.text or " ", fill=True, new_x="LMARGIN", new_y="NEXT", padding=3
            )
            pdf.ln(2)
        elif blk.kind == "table" and blk.rows:
            pdf.ln(1)
            pdf.set_font("DejaVu", "", 10)
            pdf.set_text_color(*th.ink)
            head = FontFace(emphasis="BOLD", fill_color=th.head_bg, color=th.ink)
            with pdf.table(
                headings_style=head,
                borders_layout=TableBordersLayout.HORIZONTAL_LINES,
                line_height=6,
                cell_fill_color=(252, 252, 251),
                padding=2,
            ) as table:
                for row_cells in blk.rows:
                    row = table.row()
                    for cell in row_cells:
                        row.cell(_plain(cell))
            pdf.ln(3)
        elif blk.kind == "hr":
            pdf.ln(2)
            pdf.set_draw_color(*th.rule)
            pdf.set_line_width(0.3)
            y = pdf.get_y()
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(3)

    return bytes(pdf.output())


# ---------- DOCX (python-docx, native styles, themed accent) ----------


def _docx_shade(paragraph: object, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def _docx_runs(paragraph: object, text: str, *, mono: bool = False) -> None:
    for chunk, st in _runs(text):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)  # type: ignore[attr-defined]
        run.bold = "b" in st
        run.italic = "i" in st
        if mono or "code" in st:
            run.font.name = "Consolas"


def to_docx(content: str, *, theme: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    th = resolve_theme(theme, "docx")
    accent = RGBColor(*th.accent)
    muted = RGBColor(*th.muted)
    code_hex = "".join(f"{c:02X}" for c in th.code_bg)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title_done = False
    for blk in _parse(content):
        if blk.kind == "h1" and not title_done:
            title_done = True
            doc.add_heading(_plain(blk.text), level=0)
        elif blk.kind in ("h1", "h2", "h3", "h4"):
            level = {"h1": 1, "h2": 1, "h3": 2, "h4": 3}[blk.kind]
            h = doc.add_heading(_plain(blk.text), level=level)
            if th.heading_accent:
                for run in h.runs:
                    run.font.color.rgb = accent
        elif blk.kind == "para":
            _docx_runs(doc.add_paragraph(), blk.text)
        elif blk.kind == "bullet":
            style = "List Bullet" + (f" {blk.level + 1}" if blk.level else "")
            _safe_list(doc, blk.text, style, "List Bullet")
        elif blk.kind == "number":
            style = "List Number" + (f" {blk.level + 1}" if blk.level else "")
            _safe_list(doc, blk.text, style, "List Number")
        elif blk.kind == "quote":
            para = doc.add_paragraph()
            try:
                para.style = doc.styles["Quote"]
            except KeyError:
                para.paragraph_format.left_indent = Pt(18)
            _docx_runs(para, blk.text)
        elif blk.kind == "code":
            para = doc.add_paragraph()
            _docx_shade(para, code_hex)
            run = para.add_run(blk.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif blk.kind == "table" and blk.rows:
            cols = max(len(r) for r in blk.rows)
            table = doc.add_table(rows=len(blk.rows), cols=cols)
            try:
                table.style = "Light Grid Accent 1"
            except KeyError:
                table.style = "Table Grid"
            for r, row_cells in enumerate(blk.rows):
                for c, cell in enumerate(row_cells):
                    para = table.rows[r].cells[c].paragraphs[0]
                    _docx_runs(para, cell)
                    if r == 0:
                        for run in para.runs:
                            run.bold = True
        elif blk.kind == "hr":
            sep = doc.add_paragraph()
            sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sep.add_run("• • •").font.color.rgb = muted

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _safe_list(doc: object, text: str, style: str, fallback: str) -> None:
    try:
        para = doc.add_paragraph(style=style)  # type: ignore[attr-defined]
    except KeyError:
        para = doc.add_paragraph(style=fallback)  # type: ignore[attr-defined]
    _docx_runs(para, text)


# ---------- PPTX (python-pptx, themed slides with colour bars + real tables) ----------


def _rgb(color: RGB) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor(*color)  # type: ignore[no-untyped-call]


_WHITE: RGB = (255, 255, 255)


def _pptx_rect(slide: Any, left: Any, top: Any, width: Any, height: Any, color: RGB) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _pptx_title_slide(prs: Any, th: Theme, title: str) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    w, h = prs.slide_width, prs.slide_height
    if th.cover:
        _pptx_rect(slide, 0, 0, w, h, th.accent)
        title_color = _rgb(_WHITE)
        sub_color = _rgb(th.head_bg)
    else:
        _pptx_rect(slide, 0, 0, Inches(0.4), h, th.accent)
        title_color = _rgb(th.ink)
        sub_color = _rgb(th.muted)

    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), w - Inches(2.0), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(th.title_size + 8)
    run.font.bold = True
    run.font.color.rgb = title_color

    sub = tf.add_paragraph()
    sub.alignment = PP_ALIGN.CENTER
    srun = sub.add_run()
    srun.text = "erstellt mit cmn·ai"
    srun.font.size = Pt(16)
    srun.font.color.rgb = sub_color


def _pptx_heading_bar(slide: Any, prs: Any, th: Theme, heading: str) -> None:
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    w = prs.slide_width
    _pptx_rect(slide, 0, 0, w, Inches(1.15), th.accent)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.1), w - Inches(1.2), Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = heading
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = _rgb(_WHITE)


def _pptx_content_slide(prs: Any, th: Theme, heading: str, blocks: list[_Block]) -> None:
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _pptx_heading_bar(slide, prs, th, heading)
    w, h = prs.slide_width, prs.slide_height
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), w - Inches(1.4), h - Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    first = True

    def para() -> Any:
        nonlocal first
        if first and not tf.paragraphs[0].runs:
            first = False
            return tf.paragraphs[0]
        return tf.add_paragraph()

    ink = _rgb(th.ink)
    accent = _rgb(th.accent)
    muted = _rgb(th.muted)
    for blk in blocks:
        if blk.kind in ("h3", "h4"):
            p = para()
            p.space_before = Pt(6)
            r = p.add_run()
            r.text = _plain(blk.text)
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = accent
        elif blk.kind in ("bullet", "number"):
            p = para()
            p.level = min(blk.level + 1, 4)
            marker = f"{blk.marker}. " if blk.kind == "number" else "• "
            r = p.add_run()
            r.text = marker + _plain(blk.text)
            r.font.size = Pt(18)
            r.font.color.rgb = ink
        elif blk.kind == "quote":
            p = para()
            r = p.add_run()
            r.text = _plain(blk.text)
            r.font.size = Pt(18)
            r.font.italic = True
            r.font.color.rgb = muted
        elif blk.kind == "code":
            for line in blk.text.splitlines() or [" "]:
                p = para()
                r = p.add_run()
                r.text = line
                r.font.size = Pt(14)
                r.font.name = "Consolas"
                r.font.color.rgb = ink
        else:  # para
            p = para()
            r = p.add_run()
            r.text = _plain(blk.text)
            r.font.size = Pt(18)
            r.font.color.rgb = ink


def _pptx_table_slide(prs: Any, th: Theme, heading: str, rows: list[list[str]]) -> None:
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _pptx_heading_bar(slide, prs, th, heading)
    w = prs.slide_width
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    left, top = Inches(0.7), Inches(1.6)
    width, height = w - Inches(1.4), Inches(min(0.5 * n_rows + 0.2, 5.0))
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    accent = _rgb(th.accent)
    ink = _rgb(th.ink)
    white = _rgb(_WHITE)
    for r, row_cells in enumerate(rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = _plain(row_cells[c]) if c < len(row_cells) else ""
            para = cell.text_frame.paragraphs[0]
            run = para.runs[0] if para.runs else para.add_run()
            run.font.size = Pt(14)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = accent
                run.font.bold = True
                run.font.color.rgb = white
            else:
                run.font.color.rgb = ink


def to_pptx(content: str, *, theme: str = "") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    th = resolve_theme(theme, "pptx")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blocks = _parse(content)
    title = next((b.text for b in blocks if b.kind == "h1"), "cmn·ai")
    _pptx_title_slide(prs, th, _plain(title))

    heading = _plain(title)
    buffer: list[_Block] = []
    seen_title = False

    def flush() -> None:
        nonlocal buffer
        if buffer:
            _pptx_content_slide(prs, th, heading, buffer)
            buffer = []

    for block in blocks:
        if block.kind in ("h1", "h2"):
            if not seen_title and block.text == title:
                seen_title = True
                continue
            flush()
            heading = _plain(block.text)
            continue
        if block.kind == "hr":
            continue
        if block.kind == "table" and block.rows:
            flush()
            _pptx_table_slide(prs, th, heading, block.rows)
            continue
        buffer.append(block)
    flush()

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------- XLSX (openpyxl) — opens natively in Excel and Apple Numbers ----------


def _xlsx_cell_value(text: str) -> object:
    """Store numeric-looking cells as real numbers so formulas/sorting work."""
    s = _plain(text).strip()
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return s


def _sheet_title(raw: str, used: set[str], fallback: str) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", "", _plain(raw)).strip()[:31] or fallback
    title, n = clean, 2
    while title in used:
        title = f"{clean[:28]} {n}"
        n += 1
    used.add(title)
    return title


def to_xlsx(content: str, *, theme: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    th = resolve_theme(theme, "xlsx")
    accent_hex = "".join(f"{c:02X}" for c in th.accent)
    header_fill = PatternFill(start_color=accent_hex, end_color=accent_hex, fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    blocks = _parse(content)
    wb = Workbook()
    wb.remove(wb.worksheets[0])  # drop the default sheet; we create our own
    used_titles: set[str] = set()

    # Prose (headings, paragraphs, lists, quotes) goes onto a leading overview
    # sheet; every markdown table becomes its own data sheet named after the
    # heading that precedes it.
    overview: list[tuple[str, str]] = []  # (kind, text)
    tables: list[tuple[str, list[list[str]]]] = []
    heading = ""
    for blk in blocks:
        if blk.kind in ("h1", "h2", "h3", "h4"):
            heading = blk.text
            overview.append(("head", _plain(blk.text)))
        elif blk.kind == "table" and blk.rows:
            tables.append((heading, blk.rows))
        elif blk.kind in ("para", "quote"):
            overview.append(("text", _plain(blk.text)))
        elif blk.kind in ("bullet", "number"):
            overview.append(("text", f"• {_plain(blk.text)}"))

    has_prose = any(k == "text" for k, _ in overview)
    if has_prose or not tables:
        ws = wb.create_sheet(_sheet_title("Übersicht", used_titles, "Übersicht"))
        row = 1
        for kind, text in overview:
            cell = ws.cell(row=row, column=1, value=text)
            if kind == "head":
                cell.font = Font(bold=True, size=14 if row == 1 else 12)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        ws.column_dimensions["A"].width = 90

    for i, (title, rows) in enumerate(tables, start=1):
        ws = wb.create_sheet(_sheet_title(title, used_titles, f"Tabelle {i}"))
        widths: dict[int, int] = {}
        for r, row_cells in enumerate(rows, start=1):
            for c, cell_text in enumerate(row_cells, start=1):
                cell = ws.cell(row=r, column=c, value=_xlsx_cell_value(cell_text))
                widths[c] = max(widths.get(c, 0), len(str(cell.value or "")))
                if r == 1:
                    cell.fill = header_fill
                    cell.font = header_font
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 3, 10), 60)
        ws.freeze_panes = "A2"
        if rows and len(rows) > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(len(r) for r in rows))}{len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
